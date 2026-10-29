from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "outputs" / "牛来盯盘-软件宣传文档.docx"
HERO_IMAGE = ROOT / "assets" / "2d" / "bull-window-v2.png"
FEATURE_IMAGE = ROOT / "artifacts" / "bull-preview-bullish.png"

GOLD = "D6A756"
DEEP = "15120F"
RED = "E84D3D"
GREEN = "31B878"
LIGHT = "F6F1E7"
GRAY = "666666"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=120, start=120, bottom=120, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "start", "bottom", "end", "insideH", "insideV"):
        edge_data = kwargs.get(edge)
        if not edge_data:
            continue
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        for key, value in edge_data.items():
            element.set(qn(f"w:{key}"), str(value))


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_font(run, name="Microsoft YaHei", size=None, bold=None, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_text(paragraph, text, size=10.5, bold=False, color=None):
    run = paragraph.add_run(text)
    set_font(run, size=size, bold=bold, color=color)
    return run


def add_heading(doc, text, level=1):
    paragraph = doc.add_paragraph()
    paragraph.style = doc.styles[f"Heading {level}"]
    paragraph.paragraph_format.space_before = Pt(14 if level == 1 else 8)
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run(text)
    set_font(run, size=18 if level == 1 else 13, bold=True, color=DEEP if level == 1 else GOLD)
    if level == 1:
        p_pr = paragraph._p.get_or_add_pPr()
        p_bdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "14")
        bottom.set(qn("w:space"), "6")
        bottom.set(qn("w:color"), GOLD)
        p_bdr.append(bottom)
        p_pr.append(p_bdr)
    return paragraph


def add_paragraph(doc, text, bold_prefix=None, color=None):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.line_spacing = 1.35
    paragraph.paragraph_format.space_after = Pt(5)
    if bold_prefix and text.startswith(bold_prefix):
        add_text(paragraph, bold_prefix, bold=True, color=color)
        add_text(paragraph, text[len(bold_prefix):], color=color)
    else:
        add_text(paragraph, text, color=color)
    return paragraph


def add_bullets(doc, items):
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.left_indent = Cm(0.7)
        paragraph.paragraph_format.first_line_indent = Cm(-0.35)
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.paragraph_format.line_spacing = 1.25
        add_text(paragraph, item)


def add_numbered(doc, items):
    for item in items:
        paragraph = doc.add_paragraph(style="List Number")
        paragraph.paragraph_format.left_indent = Cm(0.7)
        paragraph.paragraph_format.first_line_indent = Cm(-0.35)
        paragraph.paragraph_format.space_after = Pt(3)
        add_text(paragraph, item)


def add_feature_box(doc, title, body, accent=GOLD):
    table = doc.add_table(rows=1, cols=1)
    table.autofit = True
    cell = table.cell(0, 0)
    set_cell_shading(cell, LIGHT)
    set_cell_margins(cell, top=160, start=220, bottom=160, end=220)
    set_cell_border(cell, start={"val": "single", "sz": "24", "color": accent})
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    add_text(p, title, size=12, bold=True, color=accent)
    p2 = cell.add_paragraph()
    p2.paragraph_format.line_spacing = 1.25
    add_text(p2, body, size=10)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(paragraph, "—  ", size=9, color=GRAY)
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    set_font(run, size=9, color=GRAY)
    add_text(paragraph, "  —", size=9, color=GRAY)


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.7)
    section.bottom_margin = Cm(1.6)
    section.left_margin = Cm(1.9)
    section.right_margin = Cm(1.9)

    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)

    for idx in (1, 2, 3):
        style = doc.styles[f"Heading {idx}"]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_text(p, "牛来盯盘  |  软件宣传文档", size=8.5, color=GRAY)
    add_page_number(section.footer.paragraphs[0])


def build_cover(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(26)
    p.paragraph_format.space_after = Pt(2)
    add_text(p, "牛来盯盘", size=32, bold=True, color=GOLD)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(14)
    add_text(p, "让行情住进桌面，让每一次涨跌都有回应", size=16, bold=True, color=DEEP)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(p, "3D 桌面宠物 · A 股行情实时陪伴", size=11, color=GRAY)

    if HERO_IMAGE.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(10)
        p.add_run().add_picture(str(HERO_IMAGE), width=Cm(11.2))

    table = doc.add_table(rows=1, cols=3)
    table.autofit = False
    labels = ("实时行情", "涨跌播报", "桌面常驻")
    for cell, label in zip(table.rows[0].cells, labels):
        cell.width = Cm(5.2)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, DEEP)
        set_cell_margins(cell, top=140, start=80, bottom=140, end=80)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_text(p, label, size=11, bold=True, color=GOLD)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(18)
    add_text(p, "产品宣传与功能介绍  ·  版本 1.0.0", size=10, color=GRAY)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(p, "2026 年 8 月 18 日", size=9, color=GRAY)
    doc.add_page_break()


def build_document():
    doc = Document()
    configure_document(doc)
    build_cover(doc)

    add_heading(doc, "产品概览")
    add_paragraph(
        doc,
        "牛来盯盘是一款面向 A 股用户的轻量桌面行情伙伴。它把实时行情、持仓盈亏与一只会呼吸、会睡觉、会庆祝、也会失落的 3D 牛结合在一起，让冷冰冰的数字变成一眼就能感知的桌面反馈。",
    )
    add_paragraph(
        doc,
        "无需反复切换行情软件，抬眼看看桌面上的“牛来”，即可快速了解关注标的的价格方向、涨跌幅与持仓变化。",
    )
    add_feature_box(
        doc,
        "产品主张",
        "看得见行情，也看得见情绪。牛来盯盘，让盯盘轻一点、有趣一点。",
        accent=RED,
    )

    add_heading(doc, "核心功能")
    add_heading(doc, "1. 桌面常驻，行情抬眼可见", level=2)
    add_bullets(doc, [
        "300×300 像素透明无边框窗口，始终置顶，不遮挡主要工作区域。",
        "支持鼠标拖拽，可自由放置在屏幕任意位置。",
        "不占用任务栏，适合工作、学习或娱乐时长期陪伴。",
    ])

    add_heading(doc, "2. 行情驱动的 3D 动态角色", level=2)
    add_paragraph(doc, "牛来不是静态摆件，而是会根据 A 股交易时间和实时涨跌自动改变状态的桌面宠物。")
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    headers = ("行情状态", "牛来表现", "信息反馈")
    for cell, text in zip(table.rows[0].cells, headers):
        set_cell_shading(cell, DEEP)
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_text(p, text, size=9.5, bold=True, color="FFFFFF")
    set_repeat_table_header(table.rows[0])
    rows = [
        ("休盘", "侧卧入睡、呼吸起伏、头顶飘出 ZZZ", "显示“休盘中”，保持安静"),
        ("平盘", "站立呼吸、轻微摆动", "展示当前价格与涨跌数据"),
        ("上涨", "红金光效、昂首冲刺、庆祝动作", "红色行情卡与“牛来”语音"),
        ("下跌", "绿色光效、低头拖步、失落动作", "绿色行情卡与“妈妈”语音"),
    ]
    for row_index, row_data in enumerate(rows):
        cells = table.add_row().cells
        for cell, text in zip(cells, row_data):
            set_cell_margins(cell)
            if row_index % 2 == 0:
                set_cell_shading(cell, "FAF7F0")
            p = cell.paragraphs[0]
            add_text(p, text, size=9)

    if FEATURE_IMAGE.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(8)
        p.add_run().add_picture(str(FEATURE_IMAGE), width=Cm(7.4))
        p2 = doc.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_text(p2, "上涨状态示意：红金光效、实时行情卡与互动入口", size=8.5, color=GRAY)

    add_heading(doc, "3. 自选标的，实时跟踪", level=2)
    add_bullets(doc, [
        "默认展示上证指数，也可输入沪深股票代码切换关注标的。",
        "支持 600519、sh600519、sz000001 等常用输入方式。",
        "行情卡清晰展示名称、代码、最新价、涨跌额与涨跌幅。",
        "开盘期间自动刷新行情，免去频繁手动查询。",
    ])

    add_heading(doc, "4. 持仓盈亏，一眼看懂", level=2)
    add_paragraph(doc, "填写持仓成本价与持股数量后，软件会自动计算今日盈亏、累计盈亏及累计盈亏率。")
    add_paragraph(doc, "除了标准金额，还会把盈亏换算成鸡腿、奶茶、火锅、手机或“连续吃土”等趣味生活单位，让数字更直观，也让看盘多一点情绪价值。")

    doc.add_page_break()
    add_heading(doc, "互动与个性化")
    add_heading(doc, "有温度的互动玩法", level=2)
    add_bullets(doc, [
        "单击牛身触发欢快摇摆；快速连击触发鼻孔白烟与刨地冲锋。",
        "鼠标滚轮快速轮换上证、深成与创业板指数。",
        "“上香”开启 15 秒赛博香火祈祷。",
        "“喂大阳线”触发爱心与金币奖励；“喂大阴线”触发绿色嫌弃气体动画。",
        "涨跌幅达到 ±3% 时自动出现极端行情专属彩蛋。",
        "悬停或定时触发结合当前行情变化的股民吐槽。",
    ])

    add_heading(doc, "个性化设置", level=2)
    add_paragraph(doc, "通过右上角的毛玻璃设置面板，可调整以下内容；所有设置均保存在本机，重新启动后自动恢复。")
    add_bullets(doc, [
        "关注标的代码",
        "持仓成本价与持股数量",
        "声音开关与总音量",
        "动作幅度与模型大小",
        "特效强度",
        "行情卡显示状态",
    ])

    add_heading(doc, "适用人群与场景")
    add_bullets(doc, [
        "希望在工作时轻量关注行情的上班族。",
        "不想持续盯着复杂分时图的普通投资者。",
        "喜欢桌面宠物、动态挂件与个性化工具的用户。",
        "需要快速感知自选股方向和持仓变化的 A 股用户。",
        "想为看盘过程增加陪伴感与趣味性的投资者。",
    ])
    add_feature_box(doc, "办公陪伴", "牛来安静待在桌面角落，不打断当前工作。行情出现明显变化时，角色动作与颜色会给出直观提醒。")
    add_feature_box(doc, "轻量盯盘", "无需一直打开完整行情软件，通过牛来的状态、行情卡与持仓数据快速了解市场方向。")
    add_feature_box(doc, "情绪缓冲", "上涨时一起庆祝，下跌时用夸张但可爱的动作化解紧张感，让盯盘不再只有数字压力。")

    add_heading(doc, "产品体验与系统支持")
    add_bullets(doc, [
        "Windows 便携版无需安装，双击即可运行。",
        "单实例设计：重复启动时自动唤醒现有窗口。",
        "右上角退出按钮需按住约 1 秒，减少误触关闭。",
        "具备渲染异常记录与有限次数自动恢复能力，并在本机保存运行日志。",
        "股票、持仓、音量及特效设置保存在用户本机，无需注册账户。",
        "Windows 11 x64 已提供免安装便携版。",
        "工程已配置 macOS 10.15+ 的 Intel、Apple Silicon 与 Universal 打包支持；具体发行版本以正式发布清单为准。",
    ])

    add_heading(doc, "快速上手")
    add_numbered(doc, [
        "启动“牛来盯盘”。",
        "拖动牛来到合适的桌面位置。",
        "点击左上角行情卡或右上角设置按钮。",
        "输入希望关注的沪深股票代码。",
        "如需持仓盈亏展示，填写成本价与持股数量。",
        "保存设置，开始桌面陪伴式盯盘。",
    ])

    doc.add_page_break()
    add_heading(doc, "对外宣传文案")
    add_heading(doc, "标准版", level=2)
    add_paragraph(doc, "让行情住进桌面。牛来盯盘把 A 股实时行情变成一只有情绪的 3D 桌面宠物：休盘会睡觉，上涨会冲刺庆祝，下跌会低头拖步。支持自选股票、持仓盈亏、语音反馈和趣味互动，抬眼就能看懂行情变化。")
    add_heading(doc, "精简版", level=2)
    add_paragraph(doc, "牛来盯盘——会看行情的 3D 桌面宠物。价格涨跌、持仓盈亏、趣味互动，全部住进桌面一角。")
    add_heading(doc, "社交媒体版", level=2)
    add_paragraph(doc, "休盘睡觉，上涨狂奔，下跌低头。牛来盯盘把 A 股行情做成了一只会陪你上班的 3D 桌面牛。实时行情、持仓盈亏、语音播报、赛博上香和极端行情彩蛋，一个小窗口，装下股民一天的情绪。")

    add_heading(doc, "风险提示")
    add_feature_box(
        doc,
        "投资有风险，入市需谨慎",
        "牛来盯盘是一款行情展示与桌面陪伴工具，不构成任何投资建议、收益承诺或交易依据。行情数据可能因网络、接口或系统状态产生延迟，实际交易请以证券交易所及持牌机构提供的信息为准。",
        accent=RED,
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.core_properties.title = "牛来盯盘 · 软件宣传文档"
    doc.core_properties.subject = "产品宣传与功能介绍"
    doc.core_properties.author = "牛来盯盘"
    doc.core_properties.keywords = "牛来盯盘, 3D桌面宠物, A股, 行情, 软件宣传"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
