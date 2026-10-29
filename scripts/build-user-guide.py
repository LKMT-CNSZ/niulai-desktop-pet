from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
ARTIFACTS = ROOT / "artifacts"
OUT_DIR = ARTIFACTS / "manual-work"
OUT_PATH = OUT_DIR / "牛来盯盘-1.4.0-图文使用说明.docx"

ICON = ROOT / "build" / "icon.png"
IMG_HOME = ARTIFACTS / "bull-preview-bullish.png"
IMG_SEARCH = ARTIFACTS / "bull-preview-marketsearch.png"
IMG_PICKER = ARTIFACTS / "bull-preview-stockpicker.png"
IMG_SETTINGS = ARTIFACTS / "bull-preview-settings.png"
IMG_ROAMING = ARTIFACTS / "bull-preview-roaming.png"
IMG_FEATURES = ARTIFACTS / "bull-preview-featurepanel.png"

# compact_reference_guide preset + named Niulai brand accents.
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "1F2937"
MUTED = "667085"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
BRAND_ORANGE = "FF6B47"
BRAND_GOLD = "F6B94B"
BRAND_NAVY = "101828"
WHITE = "FFFFFF"
RED_UP = "D92D20"
GREEN_DOWN = "16803C"


def set_run_font(run, name="Calibri", east_asia="Microsoft YaHei", size=None,
                 color=None, bold=None, italic=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def shade_paragraph(paragraph, fill, left_border=None):
    ppr = paragraph._p.get_or_add_pPr()
    shd = ppr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        ppr.append(shd)
    shd.set(qn("w:fill"), fill)
    if left_border:
        pbdr = ppr.find(qn("w:pBdr"))
        if pbdr is None:
            pbdr = OxmlElement("w:pBdr")
            ppr.append(pbdr)
        left = OxmlElement("w:left")
        left.set(qn("w:val"), "single")
        left.set(qn("w:sz"), "18")
        left.set(qn("w:space"), "8")
        left.set(qn("w:color"), left_border)
        pbdr.append(left)


def set_keep(paragraph, with_next=False, together=False):
    ppr = paragraph._p.get_or_add_pPr()
    if with_next:
        ppr.append(OxmlElement("w:keepNext"))
    if together:
        ppr.append(OxmlElement("w:keepLines"))


def create_numbering(doc, kind):
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    lvl.append(start)
    numfmt = OxmlElement("w:numFmt")
    numfmt.set(qn("w:val"), "bullet" if kind == "bullet" else "decimal")
    lvl.append(numfmt)
    lvltext = OxmlElement("w:lvlText")
    lvltext.set(qn("w:val"), "•" if kind == "bullet" else "%1.")
    lvl.append(lvltext)
    suff = OxmlElement("w:suff")
    suff.set(qn("w:val"), "tab")
    lvl.append(suff)
    ppr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    ppr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "270")
    ppr.append(ind)
    lvl.append(ppr)
    rpr = OxmlElement("w:rPr")
    rfonts = OxmlElement("w:rFonts")
    rfonts.set(qn("w:ascii"), "Calibri")
    rfonts.set(qn("w:hAnsi"), "Calibri")
    rfonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    rpr.append(rfonts)
    lvl.append(rpr)
    abstract.append(lvl)
    first_num = numbering.find(qn("w:num"))
    if first_num is None:
        numbering.append(abstract)
    else:
        numbering.insert(list(numbering).index(first_num), abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def apply_num(paragraph, num_id):
    ppr = paragraph._p.get_or_add_pPr()
    numpr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    numid = OxmlElement("w:numId")
    numid.set(qn("w:val"), str(num_id))
    numpr.append(ilvl)
    numpr.append(numid)
    ppr.append(numpr)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.25


def set_cell_shading(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcpr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tcpr = cell._tc.get_or_add_tcPr()
    mar = tcpr.find(qn("w:tcMar"))
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tcpr.append(mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa, indent=120):
    table.autofit = False
    tblpr = table._tbl.tblPr
    tblw = tblpr.find(qn("w:tblW"))
    if tblw is None:
        tblw = OxmlElement("w:tblW")
        tblpr.append(tblw)
    tblw.set(qn("w:w"), str(sum(widths_dxa)))
    tblw.set(qn("w:type"), "dxa")
    tblind = tblpr.find(qn("w:tblInd"))
    if tblind is None:
        tblind = OxmlElement("w:tblInd")
        tblpr.append(tblind)
    tblind.set(qn("w:w"), str(indent))
    tblind.set(qn("w:type"), "dxa")
    layout = tblpr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tblpr.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tcpr = cell._tc.get_or_add_tcPr()
            tcw = tcpr.find(qn("w:tcW"))
            if tcw is None:
                tcw = OxmlElement("w:tcW")
                tcpr.append(tcw)
            tcw.set(qn("w:w"), str(widths_dxa[idx]))
            tcw.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def mark_repeat_header(row):
    trpr = row._tr.get_or_add_trPr()
    hdr = OxmlElement("w:tblHeader")
    hdr.set(qn("w:val"), "true")
    trpr.append(hdr)


def add_footer_page_number(section):
    footer = section.footer
    footer.distance = Inches(0.492)
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run("牛来盯盘 1.4.0  ·  使用说明     ")
    set_run_font(run, size=8.5, color=MUTED)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    txt = OxmlElement("w:t")
    txt.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    r = p.add_run()._r
    r.append(fld_begin)
    r.append(instr)
    r.append(fld_sep)
    r.append(txt)
    r.append(fld_end)


def add_header(section):
    header = section.header
    header.distance = Inches(0.492)
    p = header.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run("牛来盯盘 · 桌面看盘宠物")
    set_run_font(run, size=8.5, color=MUTED, bold=True)


def add_page_title(doc, number, title, subtitle=None):
    kicker = doc.add_paragraph()
    kicker.paragraph_format.space_before = Pt(0)
    kicker.paragraph_format.space_after = Pt(2)
    run = kicker.add_run(f"{number:02d}  USER GUIDE")
    set_run_font(run, size=9, color=BRAND_ORANGE, bold=True)
    h = doc.add_paragraph(title, style="Heading 1")
    set_keep(h, with_next=True)
    if subtitle:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(10)
        run = p.add_run(subtitle)
        set_run_font(run, size=10.5, color=MUTED)


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_run_font(r1, size=11, color=INK, bold=True)
        r2 = p.add_run(text[len(bold_lead):])
        set_run_font(r2, size=11, color=INK)
    else:
        r = p.add_run(text)
        set_run_font(r, size=11, color=INK)
    return p


def add_bullet(doc, text, bullet_num):
    p = doc.add_paragraph()
    apply_num(p, bullet_num)
    r = p.add_run(text)
    set_run_font(r, size=11, color=INK)
    return p


def add_step(doc, text, decimal_num):
    p = doc.add_paragraph()
    apply_num(p, decimal_num)
    r = p.add_run(text)
    set_run_font(r, size=11, color=INK)
    return p


def add_callout(doc, label, text, fill=LIGHT_BLUE, accent=BLUE):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.13)
    p.paragraph_format.right_indent = Inches(0.08)
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.15
    shade_paragraph(p, fill, left_border=accent)
    r = p.add_run(f"  {label}  ")
    set_run_font(r, size=10.5, color=accent, bold=True)
    r = p.add_run(text)
    set_run_font(r, size=10.5, color=INK)
    set_keep(p, together=True)
    return p


def add_image(doc, path, width, caption, alt):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(3)
    shape = p.add_run().add_picture(str(path), width=Inches(width))
    shape._inline.docPr.set("descr", alt)
    set_keep(p, with_next=True, together=True)
    c = doc.add_paragraph(style="Caption")
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = c.add_run(caption)
    set_run_font(r, size=9, color=MUTED, italic=True)
    set_keep(c, together=True)


def add_market_table(doc):
    rows = [
        ("市场", "可输入示例", "识别方式"),
        ("A 股", "600519 / sh600519", "名称、6位代码、沪深前缀"),
        ("港股", "00700 / hk00700 / 00700.HK", "名称、5位代码、HK格式"),
        ("美股", "AAPL / TSLA / AAPL.US", "公司名称或股票代码"),
    ]
    table = doc.add_table(rows=len(rows), cols=3)
    table.style = "Table Grid"
    set_table_geometry(table, [1500, 3300, 4560])
    mark_repeat_header(table.rows[0])
    for r_idx, row in enumerate(rows):
        for c_idx, value in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.text = ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx == 0 else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            rr = p.add_run(value)
            set_run_font(rr, size=9.5 if r_idx else 10, color=WHITE if r_idx == 0 else INK,
                         bold=(r_idx == 0 or c_idx == 0))
            if r_idx == 0:
                set_cell_shading(cell, BLUE)
            elif r_idx % 2 == 0:
                set_cell_shading(cell, LIGHT_GRAY)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_shortcut_table(doc):
    rows = [
        ("位置 / 操作", "结果"),
        ("左上角行情牌", "打开“我的股票”，切换已保存标的；可继续查看详情或浏览器行情"),
        ("顶部 •••", "按住并拖动整个 300×300 挂件"),
        ("顶部 ↘", "窗口归位到主屏右下角，并恢复标准大小"),
        ("右上角齿轮", "打开股票、持仓、显示、巡游、声音、性能及提醒设置"),
        ("牛身单击 / 三连点", "开心摇摆 / 暴躁冲锋与烟雾"),
        ("牛身滚轮", "在上证、深成、创业板指数之间轮换"),
        ("右上角 ×", "打开退出确认；已保存股票和设置不会丢失"),
    ]
    table = doc.add_table(rows=len(rows), cols=2)
    table.style = "Table Grid"
    set_table_geometry(table, [2700, 6660])
    mark_repeat_header(table.rows[0])
    for r_idx, row in enumerate(rows):
        for c_idx, value in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            rr = p.add_run(value)
            set_run_font(rr, size=9.4, color=WHITE if r_idx == 0 else INK,
                         bold=(r_idx == 0 or c_idx == 0))
            if r_idx == 0:
                set_cell_shading(cell, BLUE)
            elif r_idx % 2 == 0:
                set_cell_shading(cell, LIGHT_GRAY)


def setup_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    configs = {
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 14, 7),
        "Heading 3": (12, DARK_BLUE, 10, 5),
    }
    for name, (size, color, before, after) in configs.items():
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    caption = doc.styles["Caption"]
    caption.font.name = "Calibri"
    caption._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    caption.font.size = Pt(9)
    caption.font.italic = True
    caption.font.color.rgb = RGBColor.from_string(MUTED)
    caption.paragraph_format.space_before = Pt(0)
    caption.paragraph_format.space_after = Pt(8)


def build():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    for required in (ICON, IMG_HOME, IMG_SEARCH, IMG_PICKER, IMG_SETTINGS, IMG_ROAMING, IMG_FEATURES):
        if not required.exists():
            raise FileNotFoundError(required)

    doc = Document()
    setup_styles(doc)
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.start_type = WD_SECTION_START.NEW_PAGE
    add_header(section)
    add_footer_page_number(section)

    props = doc.core_properties
    props.title = "牛来盯盘 1.4.0 图文使用说明"
    props.subject = "A股、港股、美股桌面盯盘宠物产品介绍与使用教程"
    props.author = "牛来盯盘"
    props.keywords = "牛来盯盘, A股, 港股, 美股, 桌面宠物, 使用说明"
    props.comments = "适用于 Windows 11 x64 便携版 1.4.0"

    bullet_num = create_numbering(doc, "bullet")
    decimal_num = create_numbering(doc, "decimal")

    # Page 1: editorial cover.
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(26)
    p.paragraph_format.space_after = Pt(18)
    shape = p.add_run().add_picture(str(ICON), width=Inches(1.55))
    shape._inline.docPr.set("descr", "牛来盯盘应用图标：金色牛头和红色上涨箭头")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("牛来盯盘")
    set_run_font(r, size=31, color=BRAND_NAVY, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(18)
    r = p.add_run("1.4.0 图文使用说明")
    set_run_font(r, size=18, color=BRAND_ORANGE, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(34)
    r = p.add_run("A 股 · 港股 · 美股  |  桌面行情 · 自选持仓 · 3D 巡游")
    set_run_font(r, size=11.5, color=MUTED, bold=True)
    add_callout(doc, "一句话介绍", "把股票行情、持仓盈亏与一只会互动的桌面牛放在桌面角落，轻量查看，随时切换。",
                fill="FFF4EF", accent=BRAND_ORANGE)
    add_body(doc, "适用版本：1.4.0 Windows x64 便携版")
    add_body(doc, "适用系统：Windows 11 x64（无需安装，完整解压后即可运行）")
    add_body(doc, "数据说明：行情来自公开接口，持仓数据仅保存在本机。")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(38)
    r = p.add_run("让盯盘少一点紧绷，多一点陪伴。")
    set_run_font(r, size=13, color=BRAND_GOLD, bold=True)

    # Page 2: product intro.
    doc.add_page_break()
    add_page_title(doc, 1, "认识牛来盯盘", "透明、置顶、可拖动的 300×300 桌面看盘宠物")
    add_body(doc, "牛来盯盘支持 A 股、港股和美股，将行情、自选、每只股票的成本与持股数量、趣味盈亏换算和桌面互动整合在一个小窗口中。")
    add_bullet(doc, "行情牌显示名称、代码、最新价、涨跌额与涨跌幅；上涨红色、下跌绿色。", bullet_num)
    add_bullet(doc, "2D 守窗、3D 动态和自动双形态可选；重要行情可触发对应动作与特效。", bullet_num)
    add_bullet(doc, "窗口始终置顶但不占任务栏，位置会自动记忆；再次启动只唤醒同一个挂件。", bullet_num)
    add_image(doc, IMG_HOME, 3.75, "图 1  主界面：行情牌、窗口控制、3D 牛和互动按钮", "牛来盯盘上涨状态主界面截图")
    add_callout(doc, "界面提示", "左上角行情牌负责股票切换与详情；右上角齿轮负责所有设置；顶部中央的 ••• 用于拖动窗口。")

    # Page 3: install and quick start.
    doc.add_page_break()
    add_page_title(doc, 2, "启动与第一次使用", "便携版无需安装，但必须完整解压")
    add_step(doc, "打开桌面文件夹“牛来盯盘_最新版_1.4.0”。", decimal_num)
    add_step(doc, "进入“牛来盯盘-1.4.0-免安装”文件夹，双击“牛来盯盘.exe”。", decimal_num)
    add_step(doc, "若从 ZIP 使用，请先完整解压；不要直接在压缩包预览窗口中运行。", decimal_num)
    add_step(doc, "首次打开后按住顶部 ••• 拖动到喜欢的位置，软件会记住窗口坐标。", decimal_num)
    add_step(doc, "点击右上角齿轮，选择市场并搜索第一只股票；点搜索结果后保存。", decimal_num)
    add_step(doc, "如需记录持仓，展开“持仓”，填写成本价和持股数，并勾选保存到“我的股票”。", decimal_num)
    add_callout(doc, "Windows 提示", "本版未配置商业代码签名证书。若首次运行出现 SmartScreen，请核对文件来自本交付文件夹后，选择“更多信息”继续运行。",
                fill="FFF8E8", accent="7A5A00")
    h = doc.add_paragraph("开始后的三个常用入口", style="Heading 2")
    add_bullet(doc, "行情牌：切换保存股票、查看当前详情、打开浏览器完整行情。", bullet_num)
    add_bullet(doc, "齿轮：搜索股票、设置持仓、显示模式、声音、巡游与提醒。", bullet_num)
    add_bullet(doc, "互动按钮：上香、喂大阳线、喂大阴线，触发趣味动画。", bullet_num)

    # Page 4: market search.
    doc.add_page_break()
    add_page_title(doc, 3, "同时搜索 A 股、港股和美股", "市场可多选，名称和代码都能搜索")
    decimal_num = create_numbering(doc, "decimal")
    add_step(doc, "点击右上角齿轮，展开“股票与自选”。", decimal_num)
    add_step(doc, "勾选需要展示和搜索的市场，可同时勾选 A 股、港股、美股。", decimal_num)
    add_step(doc, "输入名称或代码，点击“搜索”，再点击目标搜索结果。", decimal_num)
    add_step(doc, "检查当前股票后点击底部“保存”，行情牌随即切换。", decimal_num)
    add_image(doc, IMG_SEARCH, 3.3, "图 2  三市场同时搜索：示例为搜索港股“腾讯”", "股票设置页勾选A股港股美股并搜索腾讯的截图")
    add_market_table(doc)
    add_callout(doc, "没有搜索结果？", "先确认至少勾选了一个市场；代码格式不确定时可直接输入公司名称。")

    # Page 5: saved stocks and holdings.
    doc.add_page_break()
    add_page_title(doc, 4, "保存股票与记录持仓", "每只股票独立记忆成本价和持股数量")
    add_body(doc, "选中股票后，展开设置中的“持仓”，同时填写成本价与持股数，并保持“保存到我的股票”勾选。只想关注行情时可以把两项同时留空。")
    add_bullet(doc, "今日盈亏 =（当前价 - 昨收价）× 持股数。", bullet_num)
    add_bullet(doc, "累计盈亏 =（当前价 - 成本价）× 持股数。", bullet_num)
    add_bullet(doc, "累计盈亏率 =（当前价 - 成本价）÷ 成本价 × 100%。", bullet_num)
    add_image(doc, IMG_PICKER, 3.8, "图 3  点击左上角行情牌，打开“我的股票”快速切换", "我的股票面板显示腾讯控股和贵州茅台持仓的截图")
    add_callout(doc, "数据记忆", "股票、成本价、持股数和偏好只写入当前电脑。旧版单股票持仓会自动迁移为多股票记忆。")
    add_body(doc, "“查看当前详情”可查看现价、涨跌、昨收、成交量、成交额、数据源与持仓；“浏览器看盘”会调用系统默认浏览器打开公开行情页。")

    # Page 6: settings.
    doc.add_page_break()
    add_page_title(doc, 5, "设置中心", "11 个折叠分组统一管理股票、显示、声音和性能")
    add_image(doc, IMG_SETTINGS, 3.65, "图 4  设置中心：分组折叠，底部保存 / 取消始终可见", "牛来设置中心截图")
    add_bullet(doc, "2D / 3D 模式：选择 2D 守窗、3D 动态或自动双形态，并设置行情牌是否跟随。", bullet_num)
    add_bullet(doc, "桌面巡游：控制巡游、跨屏、速度和出门频率；不影响真实行情。", bullet_num)
    add_bullet(doc, "回家与睡眠：设置演出结束后的去向，并可立即召回。", bullet_num)
    add_bullet(doc, "声音 / 动作 / 特效：分别控制原声音量、动作幅度、极端行情演出和粒子质量。", bullet_num)
    add_bullet(doc, "性能：设置省电、模型大小与质量、特效质量和最大帧率。", bullet_num)
    add_callout(doc, "保存逻辑", "“恢复默认”只修改当前表单；必须再点“保存”才生效。点“取消”不会覆盖原设置。")

    # Page 7: roaming and interaction.
    doc.add_page_break()
    add_page_title(doc, 6, "3D 桌面巡游与互动", "牛可以离开家园，在当前屏幕可用区域自然行走")
    add_image(doc, IMG_ROAMING, 3.75, "图 5  3D 牛休盘回家睡眠；巡游时会显示随身行情牌", "3D牛在透明桌面窗口中睡眠的截图")
    add_bullet(doc, "默认只在当前显示器巡游；开启跨屏后可前往相邻显示器。", bullet_num)
    add_bullet(doc, "巡游窗口默认鼠标穿透，仅当光标进入牛身或随身行情牌热区时接收操作。", bullet_num)
    add_bullet(doc, "休盘、双击牛或点击“立即召回”时，牛会平滑回到最近的家园位置。", bullet_num)
    add_bullet(doc, "单击牛身触发开心摇摆；约 1.15 秒内三连点触发冲锋和鼻孔烟雾。", bullet_num)
    add_bullet(doc, "在牛身上滚动滚轮，可轮换上证、深成和创业板指数。", bullet_num)
    add_callout(doc, "看不到完整模型？", "在设置中调小“模型大小”或点击顶部 ↘ / 设置内“归位”。多屏缩放变化后也可用归位恢复。")

    # Page 8: local features.
    doc.add_page_break()
    add_page_title(doc, 7, "提醒、情绪与安静办公", "特色功能默认关闭，按需启用，不会改变行情数字")
    add_body(doc, "在齿轮设置中打开“提醒触发器”，进入特色功能面板。可配置价格、涨跌幅、成本、盈亏、快速拉升或快速跳水等条件，最多保存 20 条提醒。")
    add_image(doc, IMG_FEATURES, 3.65, "图 6  本地特色功能：情绪记忆、收盘总结、安静办公和提醒", "本地特色功能设置面板截图")
    add_bullet(doc, "情绪记忆只改变待机、表情、走动节奏和吐槽，不会修改行情或持仓。", bullet_num)
    add_bullet(doc, "收盘总结每天最多自动显示一次，可随时关闭。", bullet_num)
    add_bullet(doc, "安静办公时间内会静音、减少特效并降低巡游频率。", bullet_num)
    add_bullet(doc, "提醒可选择仅行情牌、2D 动作或 3D 演出，并可单独决定是否播放声音。", bullet_num)
    add_callout(doc, "隐私", "所有提醒、情绪、持仓和设置均保存在本机，不上传到任何账户或云端。")

    # Page 9: reference and troubleshooting.
    doc.add_page_break()
    add_page_title(doc, 8, "快捷操作与常见问题", "遇到问题先归位、等待重连，再检查网络和设置")
    add_shortcut_table(doc)
    h = doc.add_paragraph("常见问题", style="Heading 2")
    add_body(doc, "行情一直显示“连接中”或“重连中”", bold_lead="行情一直显示“连接中”或“重连中”")
    add_body(doc, "检查网络并等待下一次 3-5 秒刷新。公开行情接口可能临时限流；软件会保留最后一次有效价格并自动重试。")
    add_body(doc, "重复双击后没有第二个窗口", bold_lead="重复双击后没有第二个窗口")
    add_body(doc, "这是单实例保护的正常表现；第二次启动会唤醒并刷新已经运行的窗口。")
    add_body(doc, "窗口消失或位置异常", bold_lead="窗口消失或位置异常")
    add_body(doc, "再次启动软件，或点击顶部 ↘ / 设置内“归位”。日志位于 C:\\Users\\当前用户名\\AppData\\Roaming\\niulai-desktop-pet\\logs\\main.log。")
    add_callout(doc, "风险提示", "公开行情可能延迟、缺失或暂时不可用；持仓盈亏是本地数学换算，仅供展示与娱乐，不构成投资建议、交易指令或收益承诺。请以券商正式交易系统为准。",
                fill="FFF4EF", accent=BRAND_ORANGE)

    doc.save(OUT_PATH)
    print(OUT_PATH)


if __name__ == "__main__":
    build()
